import pytest

from servidor import busqueda, config, creacion, db
from servidor.enrutador import Enrutador, clasificar
from servidor.proveedores import Claude, NoDisponible, Respuesta


class Falso:
    def __init__(self, nombre, gratis=True, falla=False, activo=True, modelo=None):
        self.nombre, self.gratis, self.falla, self.activo = nombre, gratis, falla, activo
        self.modelo = modelo or nombre.lower()
        self.privado = not gratis
        self.llamadas = []

    def disponible(self):
        return self.activo

    def chat(self, sistema, mensajes, max_tokens=1500, pensar=False, buscar=False):
        self.llamadas.append({"sistema": sistema, "mensajes": mensajes, "pensar": pensar})
        if self.falla:
            raise NoDisponible(f"{self.nombre} caído")
        return Respuesta(f"[{self.nombre}] respuesta", self.nombre, self.modelo, 0.0 if self.gratis else 0.01)


def cerebros(**cambios):
    base = {
        "groq": Falso("Groq", modelo="gpt-oss-20b"), "groq_pro": Falso("Groq", modelo="gpt-oss-120b"),
        "cerebras": Falso("Cerebras"), "gemini": Falso("Gemini"), "openrouter": Falso("OpenRouter"),
        "claude_barato": Falso("Claude Haiku", gratis=False), "claude": Falso("Claude", gratis=False),
        "claude_fuerte": Falso("Claude Opus", gratis=False), "grok": Falso("Grok", gratis=False, activo=False),
    }
    base.update(cambios)
    return base


@pytest.fixture(autouse=True)
def sin_internet(monkeypatch):
    monkeypatch.setattr(busqueda, "web", lambda q, n=6: [])


@pytest.mark.parametrize("texto,nivel", [
    ("hola, ¿cómo estás?", "basico"),
    ("cuéntame un chiste", "basico"),
    ("Explícame cómo funciona la fotosíntesis", "intermedio"),
    ("Dame un plan de estudio para aprender Python", "intermedio"),
    ("¿Cuáles son las últimas noticias de inteligencia artificial?", "actualidad"),
    ("¿Qué está pasando hoy con el dólar?", "actualidad"),
    ("¿A qué clientes debo llamar hoy?", "privado"),
    ("Dame ideas para mejorar Apex", "privado"),
    ("Crea una imagen de un gato astronauta", "imagen"),
    ("Hazme un logo para mi tienda", "imagen"),
    ("Redacta una carta para el banco", "documento"),
    ("Haz un documento sobre la historia de Roma", "documento"),
    ("Redacta una propuesta para Mr Papa: página web a 300 mil", "equipo_documento"),
    ("Escribe un documento importante, piénsalo bien", "equipo_documento"),
    ("¿Debería comprar acciones de Nvidia?", "equipo_inversion"),
    ("Hazme un análisis técnico de TSMC", "equipo_inversion"),
    ("Piénsalo bien: ¿qué carrera me conviene estudiar?", "avanzado"),
])
def test_clasificar(texto, nivel):
    assert clasificar(texto) == nivel


def test_basico_usa_groq_gratis():
    p = cerebros()
    r = Enrutador(p).responder("hola")
    assert r["cerebro"] == "Groq" and r["nivel"] == "basico" and r["usd"] == 0
    assert p["groq"].llamadas and not p["claude"].llamadas


def test_intermedio_usa_groq_grande_pensando():
    p = cerebros()
    Enrutador(p).responder("Explícame qué es la inflación")
    assert p["groq_pro"].llamadas[0]["pensar"] is True and not p["groq"].llamadas


def test_respaldo_gratis_si_groq_falla():
    p = cerebros(groq=Falso("Groq", falla=True))
    assert Enrutador(p).responder("hola")["cerebro"] == "Cerebras"


def test_pago_solo_si_todos_los_gratis_fallan():
    caidos = {k: Falso(k, falla=True) for k in ("groq", "groq_pro", "cerebras", "gemini", "openrouter")}
    r = Enrutador(cerebros(**caidos)).responder("hola")
    assert r["cerebro"] == "Claude Haiku" and r["aviso"]


def test_lo_privado_nunca_va_a_cerebros_gratis():
    p = cerebros(claude_barato=Falso("Claude Haiku", gratis=False, activo=False),
                 claude=Falso("Claude", gratis=False, activo=False))
    with pytest.raises(NoDisponible):
        Enrutador(p).responder("¿Qué clientes tengo pendientes en el CRM?")
    assert not any(p[k].llamadas for k in ("groq", "groq_pro", "cerebras", "gemini", "openrouter"))


def test_actualidad_sin_grok_usa_groq_con_busqueda(monkeypatch):
    monkeypatch.setattr(busqueda, "web", lambda q, n=6: [{"titulo": "Noticia", "url": "https://x.co", "texto": "dato"}])
    p = cerebros()
    r = Enrutador(p).responder("últimas noticias de Tesla")
    assert r["cerebro"] == "Groq" and r["fuentes"][0]["url"] == "https://x.co"
    assert "RESULTADOS DE BÚSQUEDA" in p["groq_pro"].llamadas[0]["sistema"]


def test_actualidad_con_grok_activo_usa_grok(monkeypatch):
    monkeypatch.setattr(busqueda, "web", lambda q, n=6: [{"titulo": "N", "url": "https://x.co", "texto": "dato"}])
    p = cerebros(grok=Falso("Grok", gratis=False))
    assert Enrutador(p).responder("tendencias de hoy en X")["cerebro"] == "Grok"


def test_documento_importante_en_equipo_sin_grok(monkeypatch):
    monkeypatch.setattr(config, "FILES_DIR", config.DATA_DIR)
    p = cerebros()
    r = Enrutador(p).responder("Redacta una propuesta formal de diseño de páginas web a 300 mil")
    assert r["nivel"] == "equipo_documento"
    assert r["cerebro"] == "Claude + Gemini"          # Claude escribe, Gemini revisa (Grok inactivo)
    assert len(p["claude"].llamadas) == 2              # borrador y versión final
    assert r["archivos"][0]["archivo"].endswith(".docx")
    assert (config.FILES_DIR / r["archivos"][0]["archivo"]).exists()


def test_documento_importante_privado_lo_revisa_claude_opus(monkeypatch):
    monkeypatch.setattr(config, "FILES_DIR", config.DATA_DIR)
    p = cerebros()
    r = Enrutador(p).responder("Redacta una propuesta formal para Apex con mis clientes")
    assert r["cerebro"] == "Claude + Claude Opus" and not p["gemini"].llamadas


def test_documento_importante_con_grok_lo_revisa_grok(monkeypatch):
    monkeypatch.setattr(config, "FILES_DIR", config.DATA_DIR)
    r = Enrutador(cerebros(grok=Falso("Grok", gratis=False))).responder("Redacta una propuesta formal de venta")
    assert r["cerebro"] == "Claude + Grok"


def test_inversion_dos_analisis_y_sintesis():
    p = cerebros(grok=Falso("Grok", gratis=False))
    r = Enrutador(p).responder("¿Debería comprar acciones de Nvidia?")
    assert r["cerebro"] == "Claude + Grok" and set(r["analisis"]) == {"Claude", "Grok"}
    assert "En qué coinciden" in p["claude"].llamadas[-1]["sistema"]


def test_inversion_sin_grok_segundo_analista_gratis():
    r = Enrutador(cerebros()).responder("¿Debería comprar acciones de Nvidia?")
    assert r["cerebro"] == "Claude + Gemini"


def test_imagen(monkeypatch):
    monkeypatch.setattr(creacion, "crear_imagen", lambda pedido, prompt=None: {"archivo": "x.png", "fuente": "Gemini"})
    r = Enrutador(cerebros()).responder("Crea una imagen de un orbe azul")
    assert r["nivel"] == "imagen" and r["archivos"][0] == {"archivo": "x.png", "tipo": "imagen"}


def test_tope_de_gasto_pausa_a_claude(monkeypatch):
    monkeypatch.setattr(config, "ANTHROPIC_API_KEY", "clave-falsa")
    monkeypatch.setattr(config, "TOPE_CLAUDE_USD", 0.05)
    db.sumar_gasto("claude", 0.06)
    with pytest.raises(NoDisponible, match="tope"):
        Claude("claude-sonnet-5-5").chat("s", [{"role": "user", "content": "hola"}])


def test_word_desde_markdown(monkeypatch):
    monkeypatch.setattr(config, "FILES_DIR", config.DATA_DIR)
    from docx import Document

    md = "# Título\n\nTexto **clave**.\n\n## Sección\n- uno\n1. paso\n\n| A | B |\n|---|---|\n| 1 | 2 |\n"
    doc = Document(str(config.FILES_DIR / creacion.guardar_docx(md, "Título")))
    assert [p.style.name for p in doc.paragraphs] == ["Title", "Normal", "Heading 2", "List Bullet", "List Number"]
    assert [[c.text for c in r.cells] for r in doc.tables[0].rows] == [["A", "B"], ["1", "2"]]
