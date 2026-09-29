"""Herramientas de negocio: CRM de ventas, buscador de prospectos, bocetos web y documentos."""
import json
import re
import threading
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo

import anthropic
import requests

import memory
import usage
from config import ANTHROPIC_API_KEY, CITY, CLAUDE_MODEL, DATA_DIR, DOCS_DIR, TIMEZONE, USER_NAME
from finance_pro import open_path

CRM_PATH = DATA_DIR / "crm.json"
ESTADOS = ["prospecto", "contactado", "interesado", "propuesta", "cliente", "perdido"]
_lock = threading.Lock()


def _today():
    return datetime.now(ZoneInfo(TIMEZONE)).date()


def _slug(t: str) -> str:
    t = re.sub(r"[^a-zA-Z0-9áéíóúñÁÉÍÓÚÑ ]", "", t).strip().lower().replace(" ", "_")
    return t[:50] or "documento"


# ---------- CRM ----------
def _crm():
    try:
        return json.loads(CRM_PATH.read_text(encoding="utf-8"))
    except (FileNotFoundError, json.JSONDecodeError):
        return []


def _crm_save(items):
    with _lock:
        CRM_PATH.write_text(json.dumps(items, ensure_ascii=False, indent=2), encoding="utf-8")


def _find(items, nombre):
    n = memory._simple(nombre)
    return next((c for c in items if n in memory._simple(c["nombre"]) or n in memory._simple(c.get("negocio", ""))), None)


def crm_agregar(nombre: str, negocio: str = "", telefono: str = "", ciudad: str = CITY,
                servicio: str = "", valor_cop: float = 0, nota: str = "", proximo_contacto: str = "") -> dict:
    items = _crm()
    if _find(items, nombre):
        return {"nota": f"{nombre} ya está en el CRM; usa crm_actualizar."}
    c = {"nombre": nombre, "negocio": negocio, "telefono": telefono, "ciudad": ciudad,
         "servicio": servicio, "valor_cop": valor_cop, "estado": "prospecto",
         "notas": [f"{_today()}: {nota}"] if nota else [],
         "proximo_contacto": proximo_contacto or str(_today() + timedelta(days=1)),
         "creado": str(_today())}
    items.append(c)
    _crm_save(items)
    return {"ok": True, "cliente": c}


def crm_actualizar(nombre: str, estado: str = "", nota: str = "", proximo_contacto: str = "",
                   valor_cop: float = None, telefono: str = "") -> dict:
    items = _crm()
    c = _find(items, nombre)
    if not c:
        return {"error": f"No encontré a {nombre} en el CRM"}
    if estado:
        c["estado"] = estado if estado in ESTADOS else c["estado"]
    if nota:
        c["notas"].append(f"{_today()}: {nota}")
    if proximo_contacto:
        c["proximo_contacto"] = proximo_contacto
    if valor_cop is not None:
        c["valor_cop"] = valor_cop
    if telefono:
        c["telefono"] = telefono
    _crm_save(items)
    return {"ok": True, "cliente": c}


def crm_ver(filtro: str = "todos") -> dict:
    """filtro: todos, hoy (a quién contactar hoy o atrasados), o un estado."""
    items = _crm()
    if filtro == "hoy":
        items = [c for c in items if c["estado"] not in ("cliente", "perdido")
                 and c.get("proximo_contacto", "9999") <= str(_today())]
    elif filtro in ESTADOS:
        items = [c for c in items if c["estado"] == filtro]
    resumen = {e: sum(1 for c in _crm() if c["estado"] == e) for e in ESTADOS}
    ganado = sum(c.get("valor_cop") or 0 for c in _crm() if c["estado"] == "cliente")
    return {"clientes": [{k: c.get(k) for k in ("nombre", "negocio", "telefono", "estado", "proximo_contacto")}
                         | {"ultima_nota": (c["notas"] or [""])[-1]} for c in items[:25]],
            "embudo": resumen, "vendido_total_cop": ganado}


# ---------- Buscador de prospectos (OpenStreetMap) ----------
CATEGORIAS = {
    "restaurantes": '["amenity"~"restaurant|fast_food|cafe|bar|ice_cream"]',
    "tiendas": '["shop"]',
    "salud": '["amenity"~"pharmacy|dentist|clinic|doctors|veterinary"]',
    "belleza": '["shop"~"hairdresser|beauty|cosmetics"]',
    "servicios": '["office"]',
    "hoteles": '["tourism"~"hotel|guest_house|hostel"]',
    "gimnasios": '["leisure"~"fitness_centre|sports_centre"]',
    "todos": '["name"]["amenity"~"restaurant|fast_food|cafe|bar|pharmacy|dentist|clinic|veterinary"]',
}


def buscar_prospectos(categoria: str = "restaurantes", ciudad: str = CITY, solo_sin_web: bool = True) -> dict:
    """Negocios registrados en OpenStreetMap; marca los que NO tienen página web."""
    sel = CATEGORIAS.get(categoria.lower(), CATEGORIAS["todos"])
    q = f"""[out:json][timeout:40];
area["name"="{ciudad}"]["boundary"="administrative"]->.a;
(node{sel}(area.a);way{sel}(area.a););
out center tags 150;"""
    r = requests.post("https://overpass-api.de/api/interpreter", data={"data": q},
                      headers={"User-Agent": "Axtra-asistente-personal"}, timeout=60)
    r.raise_for_status()
    ya = {memory._simple(c[k]) for c in _crm() for k in ("nombre", "negocio") if c.get(k)}
    out = []
    for el in r.json().get("elements", []):
        t = el.get("tags", {})
        name = t.get("name")
        if not name:
            continue
        has_web = any(k in t for k in ("website", "contact:website", "url"))
        if solo_sin_web and has_web:
            continue
        out.append({"nombre": name, "tipo": t.get("amenity") or t.get("shop") or t.get("tourism") or t.get("office", ""),
                    "telefono": t.get("phone") or t.get("contact:phone", ""),
                    "direccion": " ".join(filter(None, [t.get("addr:street"), t.get("addr:housenumber")])),
                    "tiene_web": has_web, "ya_en_crm": memory._simple(name) in ya})
    return {"ciudad": ciudad, "categoria": categoria, "encontrados": len(out), "prospectos": out[:40],
            "nota": "Datos de OpenStreetMap: puede no incluir todos los negocios. "
                    "Complementa con búsqueda web o Google Maps."}


# ---------- Documentos (propuestas, informes, cotizaciones) ----------
def guardar_documento(titulo: str, contenido: str, formato: str = "docx", abrir_ahora: bool = True) -> dict:
    """contenido en texto con títulos marcados con # (markdown sencillo)."""
    DOCS_DIR.mkdir(exist_ok=True)
    name = f"{_today()}_{_slug(titulo)}"
    if formato == "docx":
        try:
            from docx import Document
            from docx.shared import Pt

            doc = Document()
            doc.styles["Normal"].font.name = "Calibri"
            doc.styles["Normal"].font.size = Pt(11)
            doc.add_heading(titulo, level=0)
            for line in contenido.splitlines():
                s = line.strip()
                if not s:
                    continue
                m = re.match(r"^(#{1,3})\s+(.*)", s)
                if m:
                    doc.add_heading(m.group(2), level=len(m.group(1)))
                elif re.match(r"^[-*•]\s+", s):
                    doc.add_paragraph(re.sub(r"^[-*•]\s+", "", s).replace("**", ""), style="List Bullet")
                else:
                    doc.add_paragraph(s.replace("**", ""))
            path = DOCS_DIR / f"{name}.docx"
            doc.save(path)
        except ImportError:
            formato = "md"
    if formato != "docx":
        path = DOCS_DIR / f"{name}.md"
        path.write_text(f"# {titulo}\n\n{contenido}", encoding="utf-8")
    if abrir_ahora:
        open_path(path)
    return {"ok": True, "archivo": str(path)}


# ---------- Generador de bocetos web ----------
BOCETO_RULES = (
    "Crea una página web de UNA sola página, en un solo archivo HTML (CSS y JS dentro), en español, "
    "moderna, profesional y adaptada a celular, para mostrarla como BOCETO de venta a un negocio local. "
    "Incluye: encabezado con nombre y eslogan, sección de productos o servicios con precios de ejemplo "
    "realistas en pesos colombianos, sección 'Nosotros', ubicación/horario, botón de WhatsApp y un "
    "formulario de pedido o contacto que VALIDE los campos pero que NO envíe nada: al enviar debe mostrar "
    "un aviso 'Vista previa: el envío se activa al contratar el sitio'. Usa imágenes de relleno con "
    "degradados o emojis (sin enlaces a imágenes externas). Responde SOLO con el código HTML completo."
)


def _generate_boceto(negocio, tipo, ciudad, detalles, path):
    import autonomy

    try:
        client = anthropic.Anthropic(api_key=ANTHROPIC_API_KEY)
        prompt = (f"Negocio: {negocio}\nTipo: {tipo}\nCiudad: {ciudad}\nDetalles: {detalles or 'ninguno'}\n"
                  f"Hecho por: {USER_NAME} (diseño web)")
        resp = client.messages.create(model=CLAUDE_MODEL, max_tokens=12000, system=BOCETO_RULES,
                                      messages=[{"role": "user", "content": prompt}])
        usage.record(resp, "bocetos web")
        html = "".join(b.text for b in resp.content if b.type == "text")
        m = re.search(r"<!DOCTYPE html>.*</html>", html, re.S | re.I) or re.search(r"<html.*</html>", html, re.S | re.I)
        path.write_text(m.group(0) if m else html, encoding="utf-8")
        open_path(path)
        autonomy.announce(f"El boceto de {negocio} está listo y abierto en el navegador.")
    except Exception as e:
        autonomy.announce(f"No pude terminar el boceto de {negocio}: {e}")


def crear_boceto(negocio: str, tipo: str = "", ciudad: str = CITY, detalles: str = "") -> dict:
    path = DOCS_DIR / "bocetos" / f"{_slug(negocio)}.html"
    path.parent.mkdir(parents=True, exist_ok=True)
    threading.Thread(target=_generate_boceto, args=(negocio, tipo, ciudad, detalles, path), daemon=True).start()
    return {"ok": True, "nota": "Estoy diseñando el boceto en segundo plano (1 a 2 minutos). Le aviso al terminar.",
            "archivo": str(path)}


TOOL_SCHEMAS = [
    {"name": "crm_agregar", "description": "Agrega un prospecto o cliente al CRM de ventas de Felipe.",
     "input_schema": {"type": "object", "properties": {"nombre": {"type": "string"}, "negocio": {"type": "string"}, "telefono": {"type": "string"}, "ciudad": {"type": "string"}, "servicio": {"type": "string"}, "valor_cop": {"type": "number"}, "nota": {"type": "string"}, "proximo_contacto": {"type": "string", "description": "AAAA-MM-DD"}}, "required": ["nombre"]}},
    {"name": "crm_actualizar", "description": f"Actualiza un cliente del CRM: estado ({', '.join(ESTADOS)}), nota, próximo contacto (AAAA-MM-DD), valor, teléfono.",
     "input_schema": {"type": "object", "properties": {"nombre": {"type": "string"}, "estado": {"type": "string"}, "nota": {"type": "string"}, "proximo_contacto": {"type": "string"}, "valor_cop": {"type": "number"}, "telefono": {"type": "string"}}, "required": ["nombre"]}},
    {"name": "crm_ver", "description": "Ver el CRM. filtro: 'todos', 'hoy' (a quién contactar hoy) o un estado. Incluye el embudo de ventas.",
     "input_schema": {"type": "object", "properties": {"filtro": {"type": "string"}}}},
    {"name": "buscar_prospectos", "description": f"Busca negocios en una ciudad (por defecto {CITY}) que NO tienen página web, para ofrecerles bocetos. Categorías: {', '.join(CATEGORIAS)}.",
     "input_schema": {"type": "object", "properties": {"categoria": {"type": "string"}, "ciudad": {"type": "string"}, "solo_sin_web": {"type": "boolean"}}}},
    {"name": "guardar_documento", "description": "Guarda y abre un documento Word (propuestas, cotizaciones, informes). El contenido usa # para títulos y - para viñetas.",
     "input_schema": {"type": "object", "properties": {"titulo": {"type": "string"}, "contenido": {"type": "string"}, "formato": {"type": "string", "enum": ["docx", "md"]}}, "required": ["titulo", "contenido"]}},
    {"name": "crear_boceto", "description": "Diseña en segundo plano el boceto de página web para un negocio y lo abre en el navegador.",
     "input_schema": {"type": "object", "properties": {"negocio": {"type": "string"}, "tipo": {"type": "string"}, "ciudad": {"type": "string"}, "detalles": {"type": "string", "description": "colores, productos, teléfono, estilo"}}, "required": ["negocio"]}},
]
FUNCS = {"crm_agregar": crm_agregar, "crm_actualizar": crm_actualizar, "crm_ver": crm_ver,
         "buscar_prospectos": buscar_prospectos, "guardar_documento": guardar_documento,
         "crear_boceto": crear_boceto}
