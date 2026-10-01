"""Imágenes (Gemini, respaldo gratis Pollinations) y documentos Word a partir de Markdown."""
import base64
import re
import unicodedata
import urllib.parse
from datetime import datetime

import requests

from . import config
from .proveedores import GEMINI_URL, NoDisponible

_imagen = {"modelo": config.GEMINI_MODELO_IMAGEN}


def _slug(texto: str, n: int = 40) -> str:
    t = unicodedata.normalize("NFD", texto.lower())
    t = "".join(c for c in t if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]+", "_", t).strip("_")[:n] or "archivo"


def _nombre(base: str, ext: str) -> str:
    return f"{datetime.now():%Y-%m-%d_%H%M%S}_{_slug(base)}.{ext}"


# ---------------- Imágenes ----------------
def _gemini_elegir_modelo() -> str:
    r = requests.get(f"{GEMINI_URL}/models", params={"key": config.GEMINI_API_KEY, "pageSize": 200}, timeout=15)
    r.raise_for_status()
    nombres = [m["name"].split("/")[-1] for m in r.json().get("models", [])
               if "generateContent" in m.get("supportedGenerationMethods", [])]
    candidatos = [n for n in nombres if "image" in n and "imagen" not in n]
    if not candidatos:
        raise NoDisponible("tu clave de Gemini no tiene modelos de imágenes")
    return sorted(candidatos)[-1]


def _gemini(prompt: str) -> bytes:
    body = {"contents": [{"parts": [{"text": prompt}]}],
            "generationConfig": {"responseModalities": ["TEXT", "IMAGE"]}}
    for intento in range(2):
        r = requests.post(f"{GEMINI_URL}/models/{_imagen['modelo']}:generateContent",
                          headers={"x-goog-api-key": config.GEMINI_API_KEY}, json=body, timeout=120)
        if r.status_code == 404 and intento == 0:
            _imagen["modelo"] = _gemini_elegir_modelo()
            continue
        if r.status_code == 429:
            raise NoDisponible("Gemini sin cuota gratis de imágenes ahora")
        if r.status_code >= 400:
            raise NoDisponible(f"Gemini error {r.status_code}")
        for cand in r.json().get("candidates") or []:
            for part in (cand.get("content") or {}).get("parts", []):
                data = (part.get("inlineData") or part.get("inline_data") or {}).get("data")
                if data:
                    return base64.b64decode(data)
        raise NoDisponible("Gemini no devolvió imagen")
    raise NoDisponible("Gemini no respondió")


def _pollinations(prompt: str) -> bytes:
    r = requests.get("https://image.pollinations.ai/prompt/" + urllib.parse.quote(prompt[:900]),
                     params={"width": 1024, "height": 1024, "nologo": "true"}, timeout=120)
    if r.status_code >= 400 or not r.headers.get("content-type", "").startswith("image/"):
        raise NoDisponible(f"servicio gratis de imágenes no disponible ({r.status_code})")
    return r.content


def crear_imagen(pedido: str, prompt_en: str = None) -> dict:
    """Crea la imagen y la guarda. Devuelve {"archivo", "fuente"}."""
    prompt = prompt_en or pedido
    errores = []
    fuentes = ([("Gemini", _gemini)] if config.GEMINI_API_KEY else []) + [("gratis", _pollinations)]
    for nombre, fn in fuentes:
        try:
            data = fn(prompt)
        except (NoDisponible, requests.RequestException) as e:
            errores.append(f"{nombre}: {e}")
            continue
        ext = "jpg" if data[:3] == b"\xff\xd8\xff" else "png"
        archivo = _nombre(pedido, ext)
        (config.FILES_DIR / archivo).write_bytes(data)
        return {"archivo": archivo, "fuente": nombre}
    raise NoDisponible("No pude crear la imagen (" + "; ".join(errores) + ")")


# ---------------- Documentos Word ----------------
def _runs(par, texto: str) -> None:
    for i, trozo in enumerate(re.split(r"\*\*(.+?)\*\*", texto)):
        if trozo:
            par.add_run(trozo.replace("*", "")).bold = (i % 2 == 1)


def guardar_docx(md: str, titulo: str) -> str:
    """Convierte el Markdown en un Word y devuelve el nombre del archivo."""
    from docx import Document

    doc, filas, puso_titulo = Document(), [], False

    def tabla():
        rows = [r for r in filas if not re.fullmatch(r"\|?[\s:\-|]+\|?", r)]
        filas.clear()
        if not rows:
            return
        celdas = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
        t = doc.add_table(rows=len(celdas), cols=max(len(r) for r in celdas))
        t.style = "Table Grid"
        for i, fila in enumerate(celdas):
            for j, valor in enumerate(fila):
                p = t.cell(i, j).paragraphs[0]
                _runs(p, valor)
                for run in p.runs:
                    run.bold = run.bold or i == 0

    for linea in md.splitlines():
        s = linea.strip()
        if s.startswith("|"):
            filas.append(s)
            continue
        if filas:
            tabla()
        if not s or s in ("```", "---"):
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", s)
        if m:
            nivel, texto = len(m.group(1)), m.group(2).replace("*", "")
            if nivel == 1 and not puso_titulo:
                doc.add_heading(texto, 0)
                puso_titulo = True
            else:
                doc.add_heading(texto, min(nivel, 3))
        elif re.match(r"^[-*•]\s+", s):
            _runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^[-*•]\s+", "", s))
        elif re.match(r"^\d+[.)]\s+", s):
            _runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+[.)]\s+", "", s))
        else:
            _runs(doc.add_paragraph(), s)
    if filas:
        tabla()
    archivo = _nombre(titulo, "docx")
    doc.save(str(config.FILES_DIR / archivo))
    return archivo
