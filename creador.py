"""Axtra crea imágenes y documentos (lo usa la página web).

- Imágenes: Gemini (con tu GEMINI_API_KEY). Si no hay cuota o falla, Pollinations (gratis, sin clave).
  Privacidad: la descripción de la imagen se envía a ese servicio.
- Documentos Word: el texto lo escribe un cerebro (Gemini gratis, luego local o Claude) y se arma
  el .docx con títulos, listas y tablas.
Todo queda en mis_documentos/imagenes y mis_documentos/documentos.
"""
import base64
import os
import re
import unicodedata
import urllib.parse
from datetime import datetime

import requests

import llm
from config import DOCS_DIR, GEMINI_API_KEY, USER_NAME

IMAGE_MODEL = os.getenv("JARVIS_MODELO_IMAGEN", "gemini-2.5-flash-image")
IMG_DIR = DOCS_DIR / "imagenes"
DOC_DIR = DOCS_DIR / "documentos"
_gemini_image = {"model": IMAGE_MODEL}


def _slug(text: str, n: int = 40) -> str:
    t = unicodedata.normalize("NFD", text.lower())
    t = "".join(c for c in t if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]+", "_", t).strip("_")[:n] or "archivo"


def _stamp() -> str:
    return datetime.now().strftime("%Y-%m-%d_%H%M%S")


# ---------------- Imágenes ----------------
def _english_prompt(prompt: str) -> str:
    """Los generadores de imágenes entienden mejor una descripción detallada en inglés."""
    if not llm.free_ok():
        return prompt
    try:
        text, _ = llm.free_chat(
            "Convert the user's image request into ONE detailed English prompt for an image generator "
            "(subject, style, lighting, composition). Reply with the prompt only.",
            [{"role": "user", "content": prompt}], max_tokens=150, temperature=0.4)
        return text.strip().strip('"') or prompt
    except Exception:
        return prompt


def _gemini_pick_image_model() -> str:
    r = requests.get(f"{llm.GEMINI_URL}/models", params={"key": GEMINI_API_KEY, "pageSize": 200}, timeout=15)
    r.raise_for_status()
    names = [m["name"].split("/")[-1] for m in r.json().get("models", [])
             if "generateContent" in m.get("supportedGenerationMethods", [])]
    image = [n for n in names if "image" in n and "imagen" not in n]
    if not image:
        raise RuntimeError("tu clave de Gemini no tiene modelos que creen imágenes")
    return sorted(image)[-1]


def _gemini_image_bytes(prompt: str) -> bytes:
    body = {"contents": [{"parts": [{"text": prompt}]}],
            "generationConfig": {"responseModalities": ["TEXT", "IMAGE"]}}
    for attempt in range(2):
        r = requests.post(f"{llm.GEMINI_URL}/models/{_gemini_image['model']}:generateContent",
                          headers={"x-goog-api-key": GEMINI_API_KEY}, json=body, timeout=120)
        if r.status_code == 404 and attempt == 0:
            _gemini_image["model"] = _gemini_pick_image_model()
            continue
        if r.status_code == 429:
            raise RuntimeError("Gemini no tiene cuota gratis para imágenes ahora")
        r.raise_for_status()
        for cand in r.json().get("candidates") or []:
            for part in (cand.get("content") or {}).get("parts", []):
                data = (part.get("inlineData") or part.get("inline_data") or {}).get("data")
                if data:
                    return base64.b64decode(data)
        raise RuntimeError("Gemini no devolvió ninguna imagen")
    raise RuntimeError("Gemini no respondió")


def _pollinations_bytes(prompt: str) -> bytes:
    url = "https://image.pollinations.ai/prompt/" + urllib.parse.quote(prompt[:900])
    r = requests.get(url, params={"width": 1024, "height": 1024, "nologo": "true"}, timeout=120)
    r.raise_for_status()
    if not r.headers.get("content-type", "").startswith("image/"):
        raise RuntimeError("el servicio de imágenes no devolvió una imagen")
    return r.content


def crear_imagen(pedido: str) -> dict:
    """Crea una imagen a partir de una descripción. Devuelve {"ruta", "fuente"}."""
    prompt = _english_prompt(pedido)
    errores = []
    fuentes = ([("Gemini", _gemini_image_bytes)] if GEMINI_API_KEY else []) + [("Pollinations", _pollinations_bytes)]
    for nombre, fn in fuentes:
        try:
            data = fn(prompt)
        except Exception as e:
            errores.append(f"{nombre}: {str(e)[:120]}")
            print(f"  (imagen con {nombre} falló: {str(e)[:120]})")
            continue
        IMG_DIR.mkdir(parents=True, exist_ok=True)
        ext = "jpg" if data[:3] == b"\xff\xd8\xff" else "png"
        path = IMG_DIR / f"{_stamp()}_{_slug(pedido)}.{ext}"
        path.write_bytes(data)
        return {"ruta": path, "fuente": nombre}
    raise RuntimeError("No pude crear la imagen (" + "; ".join(errores) + ")")


# ---------------- Documentos Word ----------------
DOC_SYSTEM = (
    f"Eres AXTRA, el asistente personal de {USER_NAME}. Escribe el documento que te pide, completo, claro "
    "y bien organizado, en español. Formato Markdown: '# ' para el título (solo uno, al inicio), '## ' y "
    "'### ' para secciones, '- ' para viñetas, '1. ' para pasos, **negritas** para lo clave y tablas con "
    "'|' cuando ayuden. Escribe solo el documento, sin comentarios antes ni después. Nunca inventes datos "
    "que no conoces: si faltan, deja un espacio para completar entre corchetes. Si es un trabajo "
    "académico para entregar como propio, haz un esquema detallado con las ideas y fuentes sugeridas en "
    "vez del trabajo terminado."
)


def _add_runs(par, text: str) -> None:
    """Agrega texto a un párrafo respetando **negritas**."""
    for i, chunk in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if chunk:
            par.add_run(chunk.replace("*", "")).bold = (i % 2 == 1)


def markdown_a_docx(md: str, path) -> str:
    """Arma un Word a partir del Markdown. Devuelve el título."""
    from docx import Document

    doc = Document()
    title = ""
    table_rows = []

    def flush_table():
        rows = [r for r in table_rows if not re.fullmatch(r"\|?[\s:\-|]+\|?", r)]
        table_rows.clear()
        if not rows:
            return
        cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
        cols = max(len(r) for r in cells)
        table = doc.add_table(rows=len(cells), cols=cols)
        table.style = "Table Grid"
        for i, row in enumerate(cells):
            for j, value in enumerate(row):
                cell_par = table.cell(i, j).paragraphs[0]
                _add_runs(cell_par, value)
                if i == 0:
                    for run in cell_par.runs:
                        run.bold = True

    for raw in md.splitlines():
        line = raw.rstrip()
        if line.strip().startswith("|"):
            table_rows.append(line.strip())
            continue
        if table_rows:
            flush_table()
        s = line.strip()
        if not s or s in ("```", "---"):
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", s)
        if m:
            level = len(m.group(1))
            text = m.group(2).replace("*", "")
            if level == 1 and not title:
                title = text
                doc.add_heading(text, 0)
            else:
                doc.add_heading(text, min(level, 3))
        elif re.match(r"^[-*•]\s+", s):
            _add_runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^[-*•]\s+", "", s))
        elif re.match(r"^\d+[.)]\s+", s):
            _add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+[.)]\s+", "", s))
        else:
            _add_runs(doc.add_paragraph(), s)
    if table_rows:
        flush_table()
    doc.save(str(path))
    return title


def crear_documento(pedido: str) -> dict:
    """Escribe un documento y lo guarda en Word. Devuelve {"ruta", "titulo", "markdown"}."""
    md = llm.ask_text(DOC_SYSTEM, [{"role": "user", "content": pedido}], prefer="gemini",
                      max_tokens=4000, origen="documento web", temperature=0.5)
    md = re.sub(r"^```(?:markdown)?\s*|\s*```$", "", md.strip())
    DOC_DIR.mkdir(parents=True, exist_ok=True)
    first = next((l for l in md.splitlines() if l.strip()), pedido)
    path = DOC_DIR / f"{_stamp()}_{_slug(first.lstrip('# '))}.docx"
    title = markdown_a_docx(md, path) or first.lstrip("# ")
    return {"ruta": path, "titulo": title, "markdown": md}
